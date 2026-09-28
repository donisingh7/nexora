from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from typing import Protocol

from docx import Document as WordDocument
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


class UnsupportedScannedDocumentError(DocumentParseError):
    pass


@dataclass(frozen=True)
class ParsedUnit:
    text: str
    page_number: int | None = None
    location: dict[str, object] | None = None


class DocumentParser(Protocol):
    def parse(self, content: bytes) -> list[ParsedUnit]: ...


def _require_text(units: list[ParsedUnit]) -> list[ParsedUnit]:
    nonempty = [unit for unit in units if unit.text.strip()]
    if not nonempty:
        raise DocumentParseError("Document contains no extractable text")
    return nonempty


class PdfParser:
    def parse(self, content: bytes) -> list[ParsedUnit]:
        try:
            reader = PdfReader(BytesIO(content))
            units = [
                ParsedUnit(text=page.extract_text() or "", page_number=index)
                for index, page in enumerate(reader.pages, start=1)
            ]
        except Exception as exc:
            raise DocumentParseError("Unable to read PDF document") from exc
        units = [unit for unit in units if unit.text.strip()]
        if not any(any(character.isalnum() for character in unit.text) for unit in units):
            raise UnsupportedScannedDocumentError(
                "PDF has no extractable text; scanned PDFs are unsupported because "
                "OCR is not enabled"
            )
        return units


class DocxParser:
    def parse(self, content: bytes) -> list[ParsedUnit]:
        try:
            document = WordDocument(BytesIO(content))
        except Exception as exc:
            raise DocumentParseError("Unable to read DOCX document") from exc

        units: list[ParsedUnit] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.text.strip():
                units.append(ParsedUnit(paragraph.text, location={"paragraph": index}))
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    units.append(
                        ParsedUnit(
                            text,
                            location={"table": table_index, "row": row_index},
                        )
                    )
        return _require_text(units)


class TextParser:
    def parse(self, content: bytes) -> list[ParsedUnit]:
        try:
            text = content.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError as exc:
            raise DocumentParseError("Text documents must use UTF-8 encoding") from exc
        return _require_text([ParsedUnit(text, location={"source": "text"})])


class MarkdownParser(TextParser):
    pass


class ParserRegistry:
    def __init__(self) -> None:
        self._parsers: dict[str, DocumentParser] = {
            ".pdf": PdfParser(),
            ".docx": DocxParser(),
            ".txt": TextParser(),
            ".md": MarkdownParser(),
            ".markdown": MarkdownParser(),
        }

    def parse(self, filename: str, content: bytes) -> list[ParsedUnit]:
        extension = PurePath(filename).suffix.lower()
        parser = self._parsers.get(extension)
        if parser is None:
            raise DocumentParseError(f"Unsupported document format: {extension or 'unknown'}")
        return parser.parse(content)
