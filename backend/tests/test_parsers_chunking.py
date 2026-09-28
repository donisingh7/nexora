from io import BytesIO

import pytest
from docx import Document as WordDocument

from app.ingestion import parsers
from app.ingestion.chunking import TextChunker
from app.ingestion.parsers import (
    DocumentParseError,
    DocxParser,
    MarkdownParser,
    ParsedUnit,
    ParserRegistry,
    PdfParser,
    TextParser,
    UnsupportedScannedDocumentError,
)


def test_text_and_markdown_parsers_normalize_text() -> None:
    assert TextParser().parse(b"first\r\nsecond")[0].text == "first\nsecond"
    assert MarkdownParser().parse(b"# Heading")[0].text == "# Heading"


def test_docx_parser_reads_paragraphs_and_table_rows() -> None:
    document = WordDocument()
    document.add_paragraph("A useful paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    stream = BytesIO()
    document.save(stream)

    units = DocxParser().parse(stream.getvalue())

    assert units[0].text == "A useful paragraph"
    assert units[1].text == "left | right"
    assert units[1].location == {"table": 1, "row": 1}


def test_pdf_parser_extracts_pages_and_rejects_scanned_content(monkeypatch) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class Reader:
        def __init__(self, _: BytesIO) -> None:
            self.pages = [Page("page one"), Page("")]

    monkeypatch.setattr(parsers, "PdfReader", Reader)
    units = PdfParser().parse(b"pdf bytes")
    assert [(unit.page_number, unit.text) for unit in units] == [(1, "page one")]

    class ScannedReader:
        def __init__(self, _: BytesIO) -> None:
            self.pages = [Page("")]

    monkeypatch.setattr(parsers, "PdfReader", ScannedReader)
    with pytest.raises(UnsupportedScannedDocumentError, match="OCR is not enabled"):
        PdfParser().parse(b"scanned")


def test_registry_rejects_unsupported_extensions() -> None:
    with pytest.raises(DocumentParseError, match="Unsupported document format"):
        ParserRegistry().parse("slides.pptx", b"content")


def test_chunker_is_deterministic_preserves_location_and_overlap() -> None:
    chunker = TextChunker(chunk_size=20, overlap=5)
    unit = ParsedUnit("01234567890123456789ABCDEFGHIJ", page_number=3)

    chunks = chunker.chunk([unit])

    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))
    assert chunks[0].text == "01234567890123456789"
    assert chunks[1].text.startswith(chunks[0].text[-5:])
    assert all(chunk.page_number == 3 for chunk in chunks)
    assert chunks[0].location_metadata["char_start"] == 0


def test_chunker_keeps_short_documents_and_validates_overlap() -> None:
    assert [chunk.text for chunk in TextChunker(20, 4).chunk([ParsedUnit("short")])] == [
        "short"
    ]
    with pytest.raises(ValueError, match="smaller than chunk_size"):
        TextChunker(10, 10)
